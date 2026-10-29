#!/usr/bin/env python3
"""pptx_from_rendered.py 的回归：导出正确性 + 文本完整性校验 + 依赖门。

覆盖（对应第二轮复审 M1–M9）：
  M1  非法 lang 标签不崩溃、capitalize 仍按词首处理
  M2  SVG/canvas 截图透明度不重复应用，普通 img 仍写 OOXML 透明度
  M3  长文本校验无数量级退化
  M4  校验失败不发布产物、不覆盖既有文件、成功横幅在校验之后
  M5  缺失/多余/重复/顺序/页数/大小写
  M6  --skip-class 对行内、嵌套、SVG、图片、多类一致生效
  M7  display:none / visibility:hidden 的行内文字不导出
  M8  页面容器之上的祖先 opacity 也参与累乘
  M9  默认发现跳过可选依赖；专项验收显式返回 BLOCKED

依赖 python-pptx、Pillow、playwright（含可执行 chromium）。默认测试发现下缺依赖即跳过；
需要把缺失当成验收阻塞时，用 DW_PPTX_ACCEPTANCE=1 运行本模块。
"""
import base64
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import time
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "pptx_from_rendered.py")
if HERE not in sys.path:                 # 直接跑单文件时也能 import 被测模块
    sys.path.insert(0, HERE)

CSS = ("body{width:1920px;height:1080px;margin:0;font-family:Arial,sans-serif}"
       ".slide{position:relative;width:1920px;height:1080px;padding:80px;box-sizing:border-box}")

FIXTURE = """<!DOCTYPE html>
<html><head><meta charset="UTF-8"><style>""" + CSS + """
  .badge { display:inline-flex; align-items:center; gap:8px; font-size:32px; }
  .label { font-size:24px; text-transform:uppercase; letter-spacing:2px; }
  .lower { font-size:24px; text-transform:lowercase; }
  .cap { font-size:24px; text-transform:capitalize; }
  .ghost { width:600px; height:4px; background:#000; opacity:0.05; }
  .faded { font-size:28px; opacity:0.5; color:#000; }
  .nest { opacity:0.5; } .nest .inner { opacity:0.5; }
  .deep { font-size:28px; color:#000; }
  .deep-box { width:200px; height:20px; background:#000; }
  .bordered { width:200px; height:80px; background:#0000ff; border:2px solid #abcdef; opacity:0.4; }
  .page-opacity { opacity:0.8; } .block-opacity { opacity:0.7; }
  .inline-opacity { opacity:0.6; } .inner-opacity { opacity:0.5; } .deep-inline-opacity { opacity:0.4; }
  .image-parent { opacity:0.5; } .image-own { opacity:0.4; }
  .svg-parent { opacity:0.5; } .canvas-parent { opacity:0.5; }
  .case-run { color:#ff0000; }
</style></head><body>
<div class="slide">
  <div class="badge"><svg width="24" height="24" viewBox="0 0 10 10"><circle cx="5" cy="5" r="4" fill="red"/></svg>Open Source</div>
  <div class="badge"><svg width="24" height="24" viewBox="0 0 10 10"><path d="M5 2L8 7H2Z" fill="green"/></svg>+7.5 vs baseline</div>
  <div><div style="width:40px;height:40px;background:#ccc"></div>Text after block child</div>
  <p>Plain paragraph (control)</p>
  <div class="label">uppercase label</div>
  <div class="lower">MIXED Case Label</div>
  <div class="cap">capitalize me now</div>
  <p class="case-test" style="text-transform:capitalize">foo-bar l’esprit o'neill 中文测试 déjà vu ǳung</p>
  <p class="uppercase-test" style="text-transform:uppercase">straße İstanbul i ı é <span class="case-run">mixed Case</span></p>
  <p class="split-word" style="text-transform:capitalize">foo<span>bar</span></p>
  <div class="ghost"></div>
  <div class="faded">half opacity text</div>
  <div class="nest"><div class="inner"><span class="deep">nested opacity text</span><div class="deep-box"></div></div></div>
  <div class="bordered"></div>
  <div class="page-opacity"><div class="block-opacity"><p><span class="inline-opacity"><em class="inner-opacity"><b class="deep-inline-opacity">opacity chain</b></em></span></p><p>opaque neighbor</p></div></div>
  <div class="image-parent"><img class="image-own" src="pixel.png" style="width:32px;height:32px"></div>
  <div class="svg-parent"><svg width="60" height="60"><rect width="60" height="60" fill="#ff0000"/></svg></div>
  <div class="canvas-parent"><canvas width="60" height="60"></canvas></div>
  <p>visible <span style="display:none">HIDDEN-DISPLAY</span> tail</p>
  <p>visible2 <span style="visibility:hidden">HIDDEN-VIS</span> tail2</p>
  <p><span class="drop">skip inline</span> keep inline</p>
  <div class="drop"><p>skip nested</p></div>
  <div class="drop"><img src="pixel.png" style="width:24px;height:24px"></div>
  <svg class="drop" width="30" height="30"><rect width="30" height="30" fill="blue"/></svg>
  <p class="skip drop">skip multi</p>
</div>
<script>
  const c = document.querySelector('.canvas-parent canvas');
  const g = c.getContext('2d'); g.fillStyle = '#00ff00'; g.fillRect(0, 0, 60, 60);
</script>
</body></html>
"""

PNG_1PX = ("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jJ8sAAAAASUVORK5CYII=")


# ── 依赖探测（M9）──────────────────────────────────────────────

def _block_reason(pptx_ok, playwright_ok, chromium_exe):
    """把依赖状态翻译成 BLOCKED 原因；返回 None 表示可用。纯函数，便于单测。"""
    if not pptx_ok:
        return "BLOCKED：缺少 python-pptx（pip install python-pptx）"
    if not playwright_ok:
        return "BLOCKED：缺少 playwright（pip install playwright）"
    if not chromium_exe:
        return "BLOCKED：playwright 没有可用的 Chromium 路径"
    if not os.path.exists(chromium_exe) or not os.access(chromium_exe, os.X_OK):
        return f"BLOCKED：Chromium 不可执行（{chromium_exe}）"
    return None


_reason_cache = []


def deps_missing_reason():
    """真实探测：能否 import、以及 playwright 解析出的 chromium 是否可执行。"""
    if _reason_cache:
        return _reason_cache[0]
    pptx_ok = True
    try:
        import pptx  # noqa: F401
        from PIL import Image  # noqa: F401
    except Exception:
        pptx_ok = False
    playwright_ok = importlib.util.find_spec("playwright") is not None
    exe = None
    if pptx_ok and playwright_ok:
        try:
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                exe = p.chromium.executable_path
        except Exception:
            playwright_ok = False
    reason = _block_reason(pptx_ok, playwright_ok, exe)
    _reason_cache.append(reason)
    return reason


class PptxBlockReasonTest(unittest.TestCase):
    """M9：依赖状态判定本身的行为（不需要任何可选依赖即可运行）。"""

    def test_missing_pptx_is_blocked(self):
        self.assertTrue(_block_reason(False, True, "/x").startswith("BLOCKED"))

    def test_missing_playwright_is_blocked(self):
        self.assertTrue(_block_reason(True, False, None).startswith("BLOCKED"))

    def test_missing_or_non_executable_chromium_is_blocked(self):
        self.assertTrue(_block_reason(True, True, None).startswith("BLOCKED"))
        self.assertTrue(_block_reason(True, True, "/definitely/not/here").startswith("BLOCKED"))

    def test_all_present_is_not_blocked(self):
        self.assertIsNone(_block_reason(True, True, sys.executable))


@unittest.skipUnless(os.environ.get("DW_PPTX_ACCEPTANCE") == "1",
                     "PPTX 专项验收：设置 DW_PPTX_ACCEPTANCE=1 才运行（默认发现不阻塞其他测试）")
class PptxAcceptanceGateTest(unittest.TestCase):
    def test_acceptance_dependencies_are_usable(self):
        reason = deps_missing_reason()
        if reason:
            self.fail(f"PPTX 专项验收未执行，状态 BLOCKED：{reason}")


class PptxFromRenderedTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        reason = deps_missing_reason()
        if reason:
            raise unittest.SkipTest(reason)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = cls.tmp.name
        with open(os.path.join(cls.dir, "pixel.png"), "wb") as f:
            f.write(base64.b64decode(PNG_1PX))
        cls.html = os.path.join(cls.dir, "mixed_text.html")
        with open(cls.html, "w", encoding="utf-8") as f:
            f.write(FIXTURE)
        cls.out = os.path.join(cls.dir, "mixed.pptx")
        cls.dump = os.path.join(cls.dir, "pages.json")
        cls.stdout, cls.pages = cls.export(cls.html, cls.out, cls.dump)
        from pptx import Presentation
        cls.prs = Presentation(cls.out)

    @classmethod
    def tearDownClass(cls):
        if getattr(cls, "tmp", None):
            cls.tmp.cleanup()

    @classmethod
    def export(cls, html, out, dump=None, *args):
        cmd = [sys.executable, SCRIPT, html, "-o", out, "--selector", ".slide", *args]
        if dump:
            cmd.extend(["--dump-json", dump])
        r = subprocess.run(cmd, capture_output=True, text=True, cwd=HERE)
        pages = None
        if dump and os.path.exists(dump):
            with open(dump, encoding="utf-8") as f:
                pages = json.load(f)
        return r, pages

    def write_deck(self, name, body, lang=None, css=CSS):
        path = os.path.join(self.dir, name)
        attr = f' lang="{lang}"' if lang else ""
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"<!doctype html><html{attr}><head><meta charset='utf-8'>"
                    f"<style>{css}</style></head><body>{body}</body></html>")
        return path

    # ── 通用读取 ──────────────────────────────────────────────
    def text(self, prs=None):
        prs = prs or self.prs
        return " | ".join(sh.text_frame.text for s in prs.slides
                          for sh in s.shapes if sh.has_text_frame)

    def xml(self, slide_index=None, prs=None):
        prs = prs or self.prs
        slides = prs.slides if slide_index is None else [prs.slides[slide_index]]
        return "".join(sh._element.xml for s in slides for sh in s.shapes)

    def pictures(self, prs=None):
        prs = prs or self.prs
        return [sh for s in prs.slides for sh in s.shapes if sh._element.tag.endswith("}pic")]

    def pptx_with_text(self, text, name):
        from pptx import Presentation
        path = os.path.join(self.dir, name)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 914400, 914400).text_frame.text = text
        prs.save(path)
        return path

    # ── 1. 混合内容 ───────────────────────────────────────────
    def test_bare_text_mixed_with_children_preserved(self):
        t = self.text()
        for expected in ("Open Source", "+7.5 vs baseline", "Text after block child",
                         "Plain paragraph (control)"):
            self.assertIn(expected, t)

    # ── 2. text-transform ────────────────────────────────────
    def test_text_transform_matches_browser(self):
        t = self.text()
        self.assertIn("Foo-Bar L’esprit O'neill 中文测试 Déjà Vu ǲung", t)
        self.assertIn("STRASSE İSTANBUL I I É MIXED CASE", t)
        self.assertIn("mixed case label", t)
        self.assertIn("Capitalize Me Now", t)
        self.assertIn("Foobar", t)
        self.assertNotIn("uppercase label", t)

    def test_run_order_and_style_survive_transform(self):
        from pptx.dml.color import RGBColor
        runs = [run for sh in self.prs.slides[0].shapes if sh.has_text_frame
                for para in sh.text_frame.paragraphs for run in para.runs
                if run.text in ("STRASSE İSTANBUL I I É ", "MIXED CASE")]
        self.assertEqual([r.text for r in runs], ["STRASSE İSTANBUL I I É ", "MIXED CASE"])
        self.assertEqual(runs[1].font.color.rgb, RGBColor(255, 0, 0))

    # ── 3. opacity 累乘 ──────────────────────────────────────
    def test_ancestor_opacity_chain_multiplied_once(self):
        xml = self.xml()
        self.assertRegex(xml, r'<a:alpha val="6720"/>')     # 0.8×0.7×0.6×0.5×0.4
        self.assertRegex(xml, r'<a:alpha val="50000"/>')    # 0.5，只乘一次
        self.assertRegex(xml, r'<a:alpha val="5000"/>')     # 0.05
        bordered = [sh for s in self.prs.slides for sh in s.shapes if "ABCDEF" in sh._element.xml]
        self.assertEqual(len(bordered), 1)
        self.assertEqual(bordered[0]._element.xml.count('<a:alpha val="40000"/>'), 2)

    def test_opacity_does_not_leak_to_siblings(self):
        opaque = [sh for s in self.prs.slides for sh in s.shapes
                  if sh.has_text_frame and "Plain paragraph (control)" in sh.text_frame.text]
        self.assertEqual(len(opaque), 1)
        self.assertNotIn("<a:alpha", opaque[0]._element.xml)

    # ── M8: 页面容器之上的 wrapper opacity ───────────────────
    def test_wrapper_opacity_above_page_is_accumulated(self):
        html = self.write_deck(
            "wrapper.html",
            '<div style="opacity:.5"><div class="slide"><p>wrapped page text</p>'
            '<div style="width:50px;height:50px;background:#000"></div></div></div>'
            '<div class="slide"><p>opaque page text</p></div>',
            css="body{margin:0}.slide{position:relative;width:1200px;height:700px}")
        out = os.path.join(self.dir, "wrapper.pptx")
        r, _ = self.export(html, out)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        from pptx import Presentation
        prs = Presentation(out)
        self.assertIn('<a:alpha val="50000"/>', self.xml(0, prs))
        self.assertNotIn("<a:alpha", self.xml(1, prs))     # 相邻页不透明，不受影响

    # ── 4. M1: 非法 lang ─────────────────────────────────────
    def test_invalid_lang_tag_does_not_crash(self):
        for tag in ("en_US", "zh_CN", "not a tag"):
            html = self.write_deck(
                "lang.html",
                '<div class="slide"><p>plain text survives</p>'
                '<p style="text-transform:capitalize">foo bar</p></div>',
                lang=tag, css="body{margin:0}.slide{position:relative;width:1200px;height:700px}")
            out = os.path.join(self.dir, "lang.pptx")
            r, _ = self.export(html, out)
            self.assertEqual(r.returncode, 0, f"{tag}: {r.stdout[-300:]}{r.stderr[-500:]}")
            from pptx import Presentation
            text = " ".join(sh.text_frame.text for sh in Presentation(out).slides[0].shapes
                            if sh.has_text_frame)
            self.assertIn("plain text survives", text)
            # 非法标签退回默认 locale 后仍应做词首大写，而不是原样输出
            self.assertIn("Foo Bar", text, f"{tag}: {text}")

    # ── 5. M2: 图片透明度 ────────────────────────────────────
    def test_img_opacity_written_to_ooxml(self):
        faded = [sh for sh in self.pictures() if "alphaModFix" in sh._element.xml]
        self.assertEqual(len(faded), 1, [sh.name for sh in self.pictures()])
        self.assertIn('amt="20000"', faded[0]._element.xml)   # 祖先 0.5 × 自身 0.4

    def test_svg_and_canvas_screenshot_opacity_not_applied_twice(self):
        from PIL import Image
        assets = os.path.join(self.dir, "_pptx_assets")
        pngs = [n for n in os.listdir(assets) if n.endswith(".png")]
        canvas = [n for n in pngs if "canvas" in n][0]
        alpha = Image.open(os.path.join(assets, canvas)).convert("RGBA").getpixel((30, 30))[3]
        self.assertTrue(100 <= alpha <= 160, f"canvas PNG alpha={alpha}")
        # 截图已把 0.5 烤进 alpha，因此整页只能有上面那一个 alphaModFix（普通 img）
        self.assertEqual(sum("alphaModFix" in sh._element.xml for sh in self.pictures()), 1)

    # ── 6. M6: --skip-class ─────────────────────────────────
    def test_skip_class_covers_inline_nested_svg_img_and_multi_class(self):
        out = os.path.join(self.dir, "skipped.pptx")
        r, _ = self.export(self.html, out, None, "--skip-class", "drop")
        self.assertEqual(r.returncode, 0, r.stdout[-400:])
        from pptx import Presentation
        prs = Presentation(out)
        text = self.text(prs)
        for gone in ("skip inline", "skip nested", "skip multi"):
            self.assertNotIn(gone, text)
        self.assertIn("keep inline", text)
        # 跳过子树里的 svg/img 也不应再成为图片对象
        self.assertEqual(len(self.pictures(prs)), len(self.pictures()) - 2)

    # ── 7. M7: 隐藏文字 ──────────────────────────────────────
    def test_hidden_inline_text_is_not_exported(self):
        t = self.text()
        self.assertNotIn("HIDDEN-DISPLAY", t)
        self.assertNotIn("HIDDEN-VIS", t)
        self.assertIn("visible tail", t)
        self.assertIn("visible2 tail2", t)

    # ── 8. M5: 两向校验 ──────────────────────────────────────
    def test_verify_detects_duplicate_extra_and_empty_reference(self):
        import pptx_from_rendered as m
        cases = [
            # (reference, exported, 必须报差异, 差异应出现在 extra 侧)
            ("AB", "ABAB", True, True),
            ("AB", "AB EXTRA", True, True),
            ("", "UNEXPECTED", True, True),
            ("AB CD", "AB CD", False, False),
        ]
        for ref, got, want_diff, want_extra in cases:
            path = self.pptx_with_text(got, f"v{abs(hash((ref, got)))}.pptx")
            result = m.verify_text([ref], path)
            detected = bool(result["missing"][0] or result["extra"][0])
            self.assertEqual(detected, want_diff, (ref, got, result))
            self.assertEqual(bool(result["extra"][0]), want_extra, (ref, got, result))

    def test_verify_detects_reordering(self):
        import pptx_from_rendered as m
        path = self.pptx_with_text("second first", "reorder.pptx")
        result = m.verify_text(["first second"], path)
        self.assertTrue(result["extra"][0] or result["missing"][0], result)

    def test_verify_is_case_sensitive(self):
        import pptx_from_rendered as m
        path = self.pptx_with_text("UPPERCASE LABEL", "case.pptx")
        self.assertTrue(m.verify_text(["uppercase label"], path)["missing"][0])

    def test_verify_reports_page_count_mismatch_both_ways(self):
        import pptx_from_rendered as m
        one = self.pptx_with_text("first", "page1.pptx")
        from pptx import Presentation
        two = os.path.join(self.dir, "page2.pptx")
        prs = Presentation(one)
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(two)
        self.assertEqual(m.verify_text(["a", "b"], one)["page_count"],
                         {"reference": 2, "pptx": 1})
        self.assertEqual(m.verify_text(["a"], two)["page_count"],
                         {"reference": 1, "pptx": 2})

    def test_verify_keeps_pages_separate(self):
        import pptx_from_rendered as m
        from pptx import Presentation
        path = os.path.join(self.dir, "twopage.pptx")
        prs = Presentation()
        for body in ("page one text", "page two text"):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 914400, 914400).text_frame.text = body
        prs.save(path)
        result = m.verify_text(["page one text", "page two text"], path)
        self.assertEqual(result["missing"], [[], []])
        # 第 2 页文字缺失时只能出现在第 2 页，不能靠跨页拼接掩盖
        result = m.verify_text(["page one text", "page two MISSING"], path)
        self.assertEqual(result["missing"][0], [])
        self.assertTrue(result["missing"][1])

    def test_verify_detects_separator_change_at_chunk_boundary(self):
        # _tokens 按 64 字符切块；块边界上的空白增删不能让校验判通过
        import pptx_from_rendered as m
        cases = [
            ("x" * 65, "x" * 64 + " " + "x"),                    # 插入空格
            ("a" * 64 + " " + "b" * 64, "a" * 64 + "b" * 64),    # 去掉空格（两词被并）
        ]
        for ref, got in cases:
            path = self.pptx_with_text(got, f"edge{abs(hash((ref, got)))}.pptx")
            result = m.verify_text([ref], path)
            self.assertTrue(result["missing"][0] or result["extra"][0], (ref, got, result))

    def test_verify_still_normalizes_whitespace_amount(self):
        # 空白数量/类型仍然归一化，不因上一条修复变成误报
        import pptx_from_rendered as m
        path = self.pptx_with_text("alpha  beta\n\ngamma", "ws.pptx")
        self.assertEqual(m.verify_text(["alpha beta gamma"], path)["missing"], [[]])
        self.assertEqual(m.verify_text(["alpha beta gamma"], path)["extra"], [[]])

    def test_verify_normalizes_whitespace_only(self):
        import pptx_from_rendered as m
        self.assertEqual(m._norm("A  B\nC"), "A B C")
        self.assertEqual(m._norm("A-B"), "A-B")

    def test_report_missing_returns_status_and_categories(self):
        import io
        import pptx_from_rendered as m
        from contextlib import redirect_stdout
        for result, expected in (
            ({"missing": [["x"]], "extra": [[]], "reordered": [False], "page_count": None}, 1),
            ({"missing": [[]], "extra": [["y"]], "reordered": [False], "page_count": None}, 1),
            ({"missing": [[]], "extra": [[]], "reordered": [False],
              "page_count": {"reference": 2, "pptx": 1}}, 1),
            ({"missing": [[]], "extra": [[]], "reordered": [False], "page_count": None}, 0),
        ):
            buf = io.StringIO()
            with redirect_stdout(buf):
                status = m.report_missing(result)
            self.assertEqual(status, expected)
            if expected:
                self.assertIn("验收失败", buf.getvalue())

    # ── 9. M3: 长文本性能 ───────────────────────────────────
    def test_verify_text_long_document_is_not_quadratic(self):
        import pptx_from_rendered as m
        long_text = ("cell text " * 3000)[:27000]
        path = self.pptx_with_text(long_text, "long.pptx")
        started = time.perf_counter()
        result = m.verify_text([long_text], path)
        elapsed = time.perf_counter() - started
        self.assertEqual(result["missing"], [[]])
        self.assertEqual(result["extra"], [[]])
        # 修复前同条件约 11s；阈值放宽到 3s，避免依赖单次精确耗时
        self.assertLess(elapsed, 3.0, f"27k 字符校验耗时 {elapsed:.2f}s")

    # ── 10. M4: 产物发布 ─────────────────────────────────────
    def test_failed_verification_keeps_target_and_cleans_temp(self):
        target = os.path.join(self.dir, "protected.pptx")
        self.pptx_with_text("PREVIOUS GOOD", "protected.pptx")
        with open(target, "rb") as f:
            before = f.read()
        # 页面根节点的直写文字当前不会被导出：用它制造真实的校验失败
        bad = self.write_deck("bad.html", '<div class="slide">direct page text</div>')
        r, _ = self.export(bad, target)
        self.assertEqual(r.returncode, 1, r.stdout[-300:])
        with open(target, "rb") as f:
            self.assertEqual(f.read(), before)                        # 旧产物未被覆盖
        self.assertEqual([n for n in os.listdir(self.dir) if n.startswith(".pptx-tmp-")], [])
        self.assertIn("未发布", r.stdout)

    def test_success_publishes_and_prints_banner_after_verification(self):
        target = os.path.join(self.dir, "published.pptx")
        good = self.write_deck("good.html", '<div class="slide"><p>hello deck</p></div>')
        r, _ = self.export(good, target)
        self.assertEqual(r.returncode, 0, r.stdout[-300:])
        self.assertTrue(os.path.exists(target))
        self.assertEqual([n for n in os.listdir(self.dir) if n.startswith(".pptx-tmp-")], [])
        self.assertLess(r.stdout.index("没有丢字"), r.stdout.index("✅"))

    # ── 11. 端到端与模板 ────────────────────────────────────
    def test_default_export_reports_no_loss(self):
        self.assertEqual(self.stdout.returncode, 0, self.stdout.stdout[-400:])
        self.assertNotIn("❌", self.stdout.stdout)
        self.assertIn("没有丢字", self.stdout.stdout)

    def test_template_mode_keeps_matching_page_count(self):
        from pptx import Presentation
        template = os.path.join(self.dir, "template.pptx")
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(template)
        html = self.write_deck(
            "template-two.html",
            '<section class="slide"><p>template first</p></section>'
            '<section class="slide"><p>template second</p></section>',
            css="body{margin:0}.slide{position:relative;width:1200px;height:700px}")
        out = os.path.join(self.dir, "template-output.pptx")
        r, _ = self.export(html, out, None, "--template", template, "--layout", "Blank")
        self.assertEqual(r.returncode, 0, r.stdout[-300:])
        self.assertEqual(len(Presentation(out).slides), 2)
        self.assertIn("没有丢字", r.stdout)


if __name__ == "__main__":
    unittest.main()
